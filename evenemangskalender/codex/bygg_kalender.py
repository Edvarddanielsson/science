#!/usr/bin/env python3
"""bygg_kalender.py – byggskript till EVENEMANGSKALENDER (prompt v3).

Kommandon
  test      kontrollerar miljön (paket, Chromium, typsnitt) med ett fiktivt exempel
  init      skapar arbete/data.json: periodens datum, kontrolldatum, kommuner och sökfrågor
  kontroll  validerar data.json och visar täckningen (snabbt, ingen webbläsare)
  bygg      validerar data.json, bygger Excel, HTML och PDF, kör alla kontroller, skriver rapporten

Exempel
  python bygg_kalender.py test
  python bygg_kalender.py init --regioner "Kalmar län (K)" "Skåne län (S)" --period "vecka 42"
  python bygg_kalender.py kontroll arbete/data.json
  python bygg_kalender.py bygg arbete/data.json

Mall, paginering, filnamn och kontroller är fasta. Vid FEL: rätta data.json, inte skriptet.
Kräver Python 3.9+, playwright (med Chromium), openpyxl och pypdf.
Valfritt för rastrering: pypdfium2 + pillow (annars pdftoppm).
"""
from __future__ import annotations

import argparse
import datetime as dt
import difflib
import html
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path
from urllib.parse import urlsplit

VERSION = "3.0"
EJ = "ej angivet"

MANADER = ["januari", "februari", "mars", "april", "maj", "juni", "juli",
           "augusti", "september", "oktober", "november", "december"]
VECKODAGAR = ["måndag", "tisdag", "onsdag", "torsdag", "fredag", "lördag", "söndag"]

# SCB, "Län och kommuner 2026 i kodnummerordning".
LAN = [
    ("01", "Stockholms län", "Stockholm", "Upplands Väsby, Vallentuna, Österåker, Värmdö, Järfälla, Ekerö, Huddinge, Botkyrka, Salem, Haninge, Tyresö, Upplands-Bro, Nykvarn, Täby, Danderyd, Sollentuna, Stockholm, Södertälje, Nacka, Sundbyberg, Solna, Lidingö, Vaxholm, Norrtälje, Sigtuna, Nynäshamn"),
    ("03", "Uppsala län", "Uppsala", "Håbo, Älvkarleby, Knivsta, Heby, Tierp, Uppsala, Enköping, Östhammar"),
    ("04", "Södermanlands län", "Södermanland", "Vingåker, Gnesta, Nyköping, Oxelösund, Flen, Katrineholm, Eskilstuna, Strängnäs, Trosa"),
    ("05", "Östergötlands län", "Östergötland", "Ödeshög, Ydre, Kinda, Boxholm, Åtvidaberg, Finspång, Valdemarsvik, Linköping, Norrköping, Söderköping, Motala, Vadstena, Mjölby"),
    ("06", "Jönköpings län", "Jönköping", "Aneby, Gnosjö, Mullsjö, Habo, Gislaved, Vaggeryd, Jönköping, Nässjö, Värnamo, Sävsjö, Vetlanda, Eksjö, Tranås"),
    ("07", "Kronobergs län", "Kronoberg", "Uppvidinge, Lessebo, Tingsryd, Alvesta, Älmhult, Markaryd, Växjö, Ljungby"),
    ("08", "Kalmar län", "Kalmar", "Högsby, Torsås, Mörbylånga, Hultsfred, Mönsterås, Emmaboda, Kalmar, Nybro, Oskarshamn, Västervik, Vimmerby, Borgholm"),
    ("09", "Gotlands län", "Gotland", "Gotland"),
    ("10", "Blekinge län", "Blekinge", "Olofström, Karlskrona, Ronneby, Karlshamn, Sölvesborg"),
    ("12", "Skåne län", "Skåne", "Svalöv, Staffanstorp, Burlöv, Vellinge, Östra Göinge, Örkelljunga, Bjuv, Kävlinge, Lomma, Svedala, Skurup, Sjöbo, Hörby, Höör, Tomelilla, Bromölla, Osby, Perstorp, Klippan, Åstorp, Båstad, Malmö, Lund, Landskrona, Helsingborg, Höganäs, Eslöv, Ystad, Trelleborg, Kristianstad, Simrishamn, Ängelholm, Hässleholm"),
    ("13", "Hallands län", "Halland", "Hylte, Halmstad, Laholm, Falkenberg, Varberg, Kungsbacka"),
    ("14", "Västra Götalands län", "Västra Götaland", "Härryda, Partille, Öckerö, Stenungsund, Tjörn, Orust, Sotenäs, Munkedal, Tanum, Dals-Ed, Färgelanda, Ale, Lerum, Vårgårda, Bollebygd, Grästorp, Essunga, Karlsborg, Gullspång, Tranemo, Bengtsfors, Mellerud, Lilla Edet, Mark, Svenljunga, Herrljunga, Vara, Götene, Tibro, Töreboda, Göteborg, Mölndal, Kungälv, Lysekil, Uddevalla, Strömstad, Vänersborg, Trollhättan, Alingsås, Borås, Ulricehamn, Åmål, Mariestad, Lidköping, Skara, Skövde, Hjo, Tidaholm, Falköping"),
    ("17", "Värmlands län", "Värmland", "Kil, Eda, Torsby, Storfors, Hammarö, Munkfors, Forshaga, Grums, Årjäng, Sunne, Karlstad, Kristinehamn, Filipstad, Hagfors, Arvika, Säffle"),
    ("18", "Örebro län", "Örebro", "Lekeberg, Laxå, Hallsberg, Degerfors, Hällefors, Ljusnarsberg, Örebro, Kumla, Askersund, Karlskoga, Nora, Lindesberg"),
    ("19", "Västmanlands län", "Västmanland", "Skinnskatteberg, Surahammar, Kungsör, Hallstahammar, Norberg, Västerås, Sala, Fagersta, Köping, Arboga"),
    ("20", "Dalarnas län", "Dalarna", "Vansbro, Malung-Sälen, Gagnef, Leksand, Rättvik, Orsa, Älvdalen, Smedjebacken, Mora, Falun, Borlänge, Säter, Hedemora, Avesta, Ludvika"),
    ("21", "Gävleborgs län", "Gävleborg", "Ockelbo, Hofors, Ovanåker, Nordanstig, Ljusdal, Gävle, Sandviken, Söderhamn, Bollnäs, Hudiksvall"),
    ("22", "Västernorrlands län", "Västernorrland", "Ånge, Timrå, Härnösand, Sundsvall, Kramfors, Sollefteå, Örnsköldsvik"),
    ("23", "Jämtlands län", "Jämtland", "Ragunda, Bräcke, Krokom, Strömsund, Åre, Berg, Härjedalen, Östersund"),
    ("24", "Västerbottens län", "Västerbotten", "Nordmaling, Bjurholm, Vindeln, Robertsfors, Norsjö, Malå, Storuman, Sorsele, Dorotea, Vännäs, Vilhelmina, Åsele, Umeå, Lycksele, Skellefteå"),
    ("25", "Norrbottens län", "Norrbotten", "Arvidsjaur, Arjeplog, Jokkmokk, Överkalix, Kalix, Övertorneå, Pajala, Gällivare, Älvsbyn, Luleå, Piteå, Boden, Haparanda, Kiruna"),
]

TYPER = ["Mässa", "Seminarium", "Konferens", "Nätverksträff", "Frukostmöte",
         "Lunchträff", "Workshop", "Rekryteringsträff", "Öppet hus"]
A_STATUS = ("verifierad", "vantelista")
B_STATUS = ("obekraftad", "utesluten", "installt")
FORMAT = ("plats", "digitalt", "hybrid")
ANMALAN = ("kravs", "ingen", "okand")
REGLER = ("R1", "R2", "R3", "R4", "R5", "R6", "R9")
KATEGORIER = ("Arbetsmarknad", "Kompetensutveckling", "Företagande")
KALLTYPER = ("kommun", "komvux", "region", "lansstyrelse", "af", "naringsliv",
             "innovation", "larosate", "massor", "sociala", "aggregator", "sokning")
REGIONALA_KALLTYPER = {
    "region": "regionen", "lansstyrelse": "länsstyrelsen", "af": "Arbetsförmedlingen",
    "naringsliv": "Almi/NyföretagarCentrum", "larosate": "universitet/högskolor",
    "massor": "mässlistor (allamassor.se, syvarna.com/jobbmassor)",
    "sociala": "arrangörers inlägg (LinkedIn, Facebook)",
}
RESULTAT = ("kandidater", "inga", "ej_lasbar")

# Sökord, ett per sökfråga, kombinerade med regionen och periodens månad(er). Orden för
# jobbmässor bygger på en webbsökning i oktober 2026 av vad arrangörerna faktiskt skriver.
SOKORD = ["jobbmässa", "jobb- och utbildningsmässa", "rekryteringsmässa",
          "rekryteringsträff", "rekryteringsdag", "rekryteringsevent",
          "framtidsmässa", "kompetensmässa", "sommarjobbsmässa", "jobbdag",
          "\"träffa arbetsgivare\"", "\"ta med ditt CV\"", "snabbintervjuer",
          "utbildningsmässa", "komvuxmässa", "komvux öppet hus", "yrkesutbildning",
          "kostnadsfri kompetensutveckling", "starta eget", "företagarmässa"]
# Digitala jobbmässor och rekryteringsträffar tas med även utan regional koppling (R6).
SOKORD_NATIONELLT = ["digital jobbmässa", "digital rekryteringsträff"]
DIGITALT_UTAN_KOPPLING = ("Mässa", "Rekryteringsträff")

MALL_EVENEMANG = {
    "status": "verifierad | vantelista | obekraftad | utesluten | installt",
    "regel": "R1–R6 eller R9 (kvot) om status är utesluten, annars null",
    "kategori": "Arbetsmarknad | Kompetensutveckling | Företagande",
    "skal": "kort skäl (krävs för obekraftad, utesluten och installt)",
    "titel": "arrangörens titel ordagrant",
    "arrangor": "arrangör(er) fullständigt",
    "arrangor_kort": "null, eller högst 90 tecken om arrangor är längre",
    "typ": " | ".join(TYPER),
    "beskrivning": "en mening, 60–115 tecken",
    "datum": "ÅÅÅÅ-MM-DD (första dagen)",
    "datum_slut": "ÅÅÅÅ-MM-DD om flera dagar, annars null",
    "tid_start": "HH:MM, eller null om arrangören inte anger starttid (kortet visar 'Tid ej angiven')",
    "tid_slut": "HH:MM eller null",
    "format": "plats | digitalt | hybrid",
    "kommun": "kommunen där evenemanget äger rum (null om digitalt)",
    "ort": "ort",
    "lokal": "lokal eller null",
    "adress": "gatuadress utan postnummer, eller null",
    "preliminar_lokal": "true om arrangören kallar lokalen preliminär, annars false",
    "digital_plattform": "t.ex. Teams eller Zoom, eller null",
    "regional_koppling": "krävs om digitalt (utom jobbmässor och rekryteringsträffar): arrangör, målgrupp eller finansiering i regionen",
    "dropin": "HH:MM–HH:MM eller null",
    "malgrupp": "målgrupp eller null",
    "deltagarkrav": "deltagarkrav eller null",
    "pris": "prisuppgiften som källan anger, eller null",
    "pris_tagg": "Gratis | Fri entré | 150 kr | 150 kr medlem / 250 kr övriga ex moms | null",
    "villkor": "lista med taggar ur villkorsvokabulären, t.ex. [\"Begränsat antal platser\"]",
    "ovrigt": "högst en kort tagg, eller null",
    "anmalan": "kravs | ingen | okand",
    "anmalan_text": "källans uppgift om anmälan, eller null",
    "tidsfrist": "ÅÅÅÅ-MM-DD eller null",
    "lank": "https://… direktlänk till evenemangssidan",
    "samlingssida": "true om länken går till en samlingssida, annars false",
    "kalla": "var kandidaten hittades och var den verifierades",
    "belagg": "ordagrant citat från arrangörens sida eller inlägg med datum, år och starttid (det som anges)",
    "uppdaterad": "null, eller ÅÅÅÅ-MM-DD då sidan eller inlägget publicerades eller uppdaterades (krävs om året inte står i belagg)",
    "osakerheter": "avvikelser och tveksamheter, eller null",
    "kontrolldatum": "null (= det globala) eller ÅÅÅÅ-MM-DD",
}
MALL_KALLA = {
    "typ": " | ".join(KALLTYPER),
    "kommun": "kommunnamn, eller null för regional källa",
    "url": "https://… (ej för sokning)",
    "fraga": "sökfrågan ordagrant (endast sokning)",
    "resultat": "kandidater | inga | ej_lasbar",
    "notering": "valfri kommentar, eller null",
}

# ---------------------------------------------------------------- mall (BILAGA A i v2, oförändrad)

CSS = """:root{--navy:#143460;--orange:#C25700;--card:#f7fbfc;--line:#d7e0e7;--tag:#e8edf1}
*{box-sizing:border-box;margin:0;padding:0}
html{-webkit-text-size-adjust:100%}
body{font-family:Arial,"Liberation Sans",Helvetica,sans-serif;color:var(--navy);background:#477780;-webkit-print-color-adjust:exact;print-color-adjust:exact}
.sheet{width:210mm;height:297mm;padding:10mm;margin:0 auto 4mm;background:linear-gradient(135deg,#517fa6,#477780);position:relative;overflow:hidden;break-after:page;page-break-after:always}
.sheet:last-child{break-after:auto;page-break-after:auto;margin-bottom:0}
h1{font-size:32.5pt;font-weight:700;color:#fff;text-align:center;line-height:1.15;white-space:nowrap}
.sub{font-size:22.5pt;font-weight:400;color:#fff;text-align:center;line-height:1.2;margin-top:4mm;margin-bottom:7mm}
.sub b{font-weight:900}
.day{margin-top:8mm}
.day.first{margin-top:0}
.dh{margin-bottom:2mm;break-after:avoid}
.dh span{display:inline-block;background:var(--orange);color:#fff;font-size:16pt;font-weight:700;line-height:1.2;padding:1.4mm 3mm;border-radius:2mm}
a.card{display:block;background:var(--card);border-radius:4mm;padding:6mm;color:var(--navy);text-decoration:none;margin-bottom:2mm;break-inside:avoid}
a.card:last-child{margin-bottom:0}
a.card:hover,a.card:focus-visible{background:#e9f2f7}
a.card:active{background:#d9e8f0}
a.card:focus-visible{outline:3px solid var(--orange);outline-offset:1px}
.top{display:grid;grid-template-columns:2fr 1fr;column-gap:4mm}
.t{font-size:20pt;font-weight:700;line-height:1.15}
.o{font-size:12.5pt;margin-top:1.2mm;line-height:1.25}
.r{border-left:1px solid var(--line);padding-left:4mm}
.tm,.pl{display:flex;align-items:flex-start;gap:1.6mm}
.tm{font-size:15pt;font-weight:700;line-height:1.2;white-space:nowrap}
.pl{font-size:12.5pt;line-height:1.25;margin-top:1.6mm}
.tm svg{width:16.5pt;height:16.5pt;flex:none;margin-top:.2mm}
.pl svg{width:13.75pt;height:13.75pt;flex:none;margin-top:.2mm}
svg{fill:none;stroke:currentColor;stroke-width:1.7;stroke-linecap:round;stroke-linejoin:round}
.d{font-size:14pt;line-height:1.4;margin-top:3mm}
.tags{display:flex;flex-wrap:wrap;gap:1.6mm;border-top:1px solid var(--line);margin-top:3mm;padding-top:3mm}
.tag{font-size:12.5pt;line-height:1.2;padding:1mm 2.8mm;border-radius:999px;background:var(--tag);color:var(--navy)}
.tag.dl{background:#fff;color:var(--orange);border:1px solid var(--orange);padding:calc(1mm - 1px) calc(2.8mm - 1px)}
@page{size:A4;margin:0}
@media print{body{background:none}.sheet{margin:0}a.card{background:var(--card)!important}}"""

SVG_TID = ('<svg viewBox="0 0 24 24" aria-hidden="true"><circle cx="12" cy="12" r="9"/>'
           '<path d="M12 7v5l3.2 2"/></svg>')
SVG_PLATS = ('<svg viewBox="0 0 24 24" aria-hidden="true"><path d="M12 21s-7-6.2-7-11.5a7 7 0 0 1 '
             '14 0C19 14.8 12 21 12 21z"/><circle cx="12" cy="9.5" r="2.6"/></svg>')
ZOOM_JS = ('<script>(function(){var w=210/25.4*96;function f(){var s=Math.min(1,document.documentElement'
           '.clientWidth/w);document.body.style.zoom=s<1?s:""}if(!matchMedia("print").matches){f();'
           'addEventListener("resize",f)}})();</script>')

# Paginering (v2 avsnitt 6), i mm.
SIDHOJD = 277.0          # 297 mm − 2 × 10 mm padding
SAKERHET = 0.8
MARG_RUBRIK = 4.0        # h1 → underrad
MARG_UNDERRAD = 7.0      # underrad → första dagen
MARG_DAG = 8.0           # före en dag som inte står överst
MARG_DH = 2.0            # dagsrubrik → kort
MARG_KORT = 2.0          # mellan kort
RAD_D = 14 * 4 / 3 * 1.4  # radhöjd beskrivning i px (14 pt × 1,4)
RAD_H1 = 32.5 * 4 / 3 * 1.15

# ---------------------------------------------------------------- små hjälpare


def fel_avbryt(msg: str, kod: int = 2):
    print(f"FEL: {msg}", file=sys.stderr)
    sys.exit(kod)


def idag_sthlm() -> dt.date:
    try:
        from zoneinfo import ZoneInfo
        return dt.datetime.now(ZoneInfo("Europe/Stockholm")).date()
    except Exception:
        return dt.date.today()


_SV = str.maketrans({"å": "{", "ä": "|", "ö": "}", "æ": "|", "ø": "}", "ü": "y"})


def sv_nyckel(s: str) -> str:
    """Sorteringsnyckel A–Ö (å, ä, ö sist, accenter ignoreras)."""
    s = unicodedata.normalize("NFD", s.casefold().translate(_SV))
    return "".join(c for c in s if not unicodedata.combining(c))


def ascii_vikning(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    return "".join(c for c in s if not unicodedata.combining(c)).casefold()


def txt(v):
    """Trimmad text, eller None för saknat värde."""
    if v is None or isinstance(v, bool):
        return None
    if not isinstance(v, str):
        v = str(v)
    v = re.sub(r"\s+", " ", v.replace(" ", " ")).strip()
    if v == "" or v.casefold() in ("ej angivet", "null", "none", "-", "–"):
        return None
    return v


def las_datum(v):
    """date, None (saknas) eller False (ogiltigt)."""
    s = txt(v)
    if s is None:
        return None
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        return False
    try:
        return dt.date.fromisoformat(s)
    except ValueError:
        return False


def las_tid(v):
    s = txt(v)
    if s is None:
        return None
    m = re.fullmatch(r"(?:kl\.?\s*)?(\d{1,2})[:.](\d{2})", s, re.I)
    if not m or int(m.group(1)) > 23 or int(m.group(2)) > 59:
        return False
    return f"{int(m.group(1)):02d}:{m.group(2)}"


def las_intervall(v):
    s = txt(v)
    if s is None:
        return None
    delar = re.split(r"\s*[–—-]\s*", s)
    if len(delar) != 2:
        return False
    a, b = las_tid(delar[0]), las_tid(delar[1])
    if not a or not b or b <= a:
        return False
    return f"{a}–{b}"


def las_bool(v) -> bool:
    if isinstance(v, bool):
        return v
    return str(v).strip().casefold() in ("true", "ja", "1", "yes")


def tusental(s: str) -> str:
    return re.sub(r"\d{4,}", lambda m: f"{int(m.group()):,}".replace(",", " "), s)


def datum_text(d: dt.date, med_ar: bool = False) -> str:
    return f"{d.day} {MANADER[d.month - 1]}" + (f" {d.year}" if med_ar else "")


def intervall(start: dt.date, slut: dt.date, med_ar: bool = True) -> str:
    """12–18 oktober 2026 · 28 september–4 oktober 2026 · 28 december 2026–3 januari 2027"""
    if start == slut:
        return datum_text(start, med_ar)
    if start.year != slut.year:
        return f"{datum_text(start, True)}–{datum_text(slut, True)}"
    ar = f" {slut.year}" if med_ar else ""
    if start.month != slut.month:
        return f"{datum_text(start)}–{datum_text(slut)}{ar}"
    return f"{start.day}–{slut.day} {MANADER[slut.month - 1]}{ar}"


def dagrubrik(d: dt.date) -> str:
    return f"{VECKODAGAR[d.weekday()].capitalize()} {d.day} {MANADER[d.month - 1]}"


def kortdatum(d: dt.date) -> str:
    return f"{d.day} {MANADER[d.month - 1][:3]}"


def filnamn(s: str) -> str:
    return re.sub(r"[\s/\\:*?\"<>|]+", "_", s.strip())


def norm_jmf(s: str) -> str:
    return re.sub(r"[\W_]+", "", s.casefold())


def esc(s: str) -> str:
    return html.escape(s, quote=True)


def har_emoji(s: str) -> bool:
    return any(ord(c) >= 0x1F000 or 0x2600 <= ord(c) <= 0x27BF for c in s)

# ---------------------------------------------------------------- period och regioner


def berakna_period(typ: str, nr: int, ar: int):
    if typ == "vecka":
        start = dt.date.fromisocalendar(ar, nr, 1)
        return start, start + dt.timedelta(days=6)
    start = dt.date(ar, nr, 1)
    nasta = dt.date(ar + (nr == 12), nr % 12 + 1, 1)
    return start, nasta - dt.timedelta(days=1)


def tolka_period(text: str):
    t = text.strip().casefold()
    m = re.fullmatch(r"(?:vecka|v\.?)\s*(\d{1,2})", t)
    if m and 1 <= int(m.group(1)) <= 53:
        return "vecka", int(m.group(1))
    m = re.fullmatch(r"(?:månad\s+)?([a-zåäö]+)", t)
    if m and m.group(1) in MANADER:
        return "manad", MANADER.index(m.group(1)) + 1
    m = re.fullmatch(r"månad\s+(\d{1,2})", t)
    if m and 1 <= int(m.group(1)) <= 12:
        return "manad", int(m.group(1))
    fel_avbryt(f"okänd period '{text}' – skriv t.ex. \"vecka 42\" eller \"månad oktober\"")


def auto_ar(typ: str, nr: int, idag: dt.date) -> int:
    """Närmaste förekomst av perioden som inte helt har passerat."""
    bas = idag.isocalendar()[0] if typ == "vecka" else idag.year
    for ar in (bas, bas + 1, bas + 2):
        try:
            _, slut = berakna_period(typ, nr, ar)
        except ValueError:          # vecka 53 finns inte alla år
            continue
        if slut >= idag:
            return ar
    fel_avbryt(f"hittar ingen kommande förekomst av {typ} {nr}")


def manadsord(start: dt.date, slut: dt.date) -> str:
    par, d = [], start.replace(day=1)
    while d <= slut:
        par.append((d.month, d.year))
        d = (d + dt.timedelta(days=32)).replace(day=1)
    if len({y for _, y in par}) == 1:
        return " ".join(MANADER[m - 1] for m, _ in par) + f" {par[0][1]}"
    return " ".join(f"{MANADER[m - 1]} {y}" for m, y in par)


def etiketter(p: dict) -> dict:
    if p["typ"] == "vecka":
        return {"h1": f"EVENEMANG · VECKA {p['nr']}", "prefix": f"V{p['nr']}_{p['ar']}",
                "namn": f"vecka {p['nr']}", "titel": f"VECKA {p['nr']}"}
    m = MANADER[p["nr"] - 1]
    return {"h1": f"EVENEMANG · {m.upper()}", "prefix": f"{m.capitalize()}_{p['ar']}",
            "namn": m, "titel": m.upper()}


def _lan_nyckel(s: str) -> str:
    s = re.sub(r"\s+", " ", s.strip().casefold())
    s = re.sub(r"^region\s+", "", s)
    return re.sub(r"\s+län$", "", s)


LAN_INDEX = {}
for _kod, _namn, _kort, _kommuner in LAN:
    _post = {"lanskod": _kod, "namn": _namn, "kort": _kort,
             "kommuner": [k.strip() for k in _kommuner.split(",")]}
    for _alias in (_namn, _kort, _kort + "s"):
        LAN_INDEX[_lan_nyckel(_alias)] = _post
LAN_INDEX["västra götalandsregionen"] = LAN_INDEX["västra götaland"]


def tolka_region(arg: str) -> dict:
    m = re.fullmatch(r"\s*(.+?)\s*(?:\(\s*([A-Za-zÅÄÖåäö]{1,3})\s*\)|:\s*([A-Za-zÅÄÖåäö]{1,3}))?\s*", arg)
    namn, kod = m.group(1), (m.group(2) or m.group(3) or "").upper() or None
    lan = LAN_INDEX.get(_lan_nyckel(namn))
    if lan:
        return {"namn": lan["namn"], "kort": lan["kort"], "kod": kod,
                "kommuner": sorted(lan["kommuner"], key=sv_nyckel)}
    print(f"OBS: '{namn}' är inget län. Fyll i 'kommuner' i data.json själv "
          f"(t.ex. kommunerna i landskapet).")
    return {"namn": namn, "kort": namn, "kod": kod, "kommuner": []}


def tilldela_koder(regioner: list):
    def forslag(kort):
        ord_ = kort.split()
        return "".join(o[0] for o in ord_[:3]).upper() if len(ord_) > 1 else kort[:2].upper()
    for r in regioner:
        if not r["kod"]:
            r["kod"] = r["kort"][0].upper()
    koder = [r["kod"] for r in regioner]
    for r in regioner:
        if koder.count(r["kod"]) > 1:
            r["kod"] = forslag(r["kort"])
    koder = [r["kod"] for r in regioner]
    if len(set(koder)) != len(koder):
        fel_avbryt("regionernas ID-prefix krockar – ange dem själv, t.ex. \"Västra Götalands län (VG)\"")

# ---------------------------------------------------------------- validering


class Ctx:
    def __init__(self, tillat_exempel: bool = False):
        self.fel: list[str] = []
        self.varn: list[str] = []
        self.tillat_exempel = tillat_exempel


FORKORTNINGAR = ("t.ex.", "bl.a.", "m.m.", "s.k.", "inkl.", "exkl.", "ca.", "kl.", "o.s.v.",
                 "osv.", "dvs.", "t.o.m.", "fr.o.m.", "jfr.", "resp.", "nr.", "St.", "S:t")
BELOPP = r"\d{1,3}(?: \d{3})*"
PRIS_RE = [re.compile(p) for p in (
    r"Gratis", r"Fri entré", rf"{BELOPP} kr(?: ex moms)?",
    rf"{BELOPP} kr medlem / {BELOPP} kr övriga(?: ex moms)?")]
VILLKOR_RE = [re.compile(p) for p in (
    r"Endast \S.{1,38}", r"Från \d{1,2} år", r"För \S.{1,38}", r"Begränsat antal platser",
    r"Väntelista", r"Bindande anmälan", rf"Uteblivande {BELOPP} kr(?: ex moms)?")]


def kontrollera_beskrivning(b: str, titel: str) -> tuple[list, list]:
    fel, varn = [], []
    if not 60 <= len(b) <= 115:
        fel.append(f"beskrivning har {len(b)} tecken (ska vara 60–115)")
    if not b.endswith("."):
        fel.append("beskrivning ska vara en mening som slutar med punkt")
    s = b
    for f in FORKORTNINGAR:
        s = s.replace(f, f.replace(".", ""))
    if re.search(r"[.!?]\s+\S", s[:-1]):
        fel.append("beskrivning ska vara en enda mening")
    if re.search(r"\b([01]?\d|2[0-3])[:.][0-5]\d\b|\bkl\.?\s*\d", b, re.I):
        fel.append("beskrivning får inte upprepa tiden")
    if re.search(r"\d\s?(kr|kronor)\b", b, re.I):
        fel.append("beskrivning får inte upprepa priset")
    elif re.search(r"\b(gratis|kostnadsfri\w*|avgiftsfri\w*)\b", b, re.I):
        varn.append("beskrivning nämner pris (gratis/kostnadsfritt) – upprepa inte priset")
    if len(titel) >= 12 and norm_jmf(titel) in norm_jmf(b):
        fel.append("beskrivning får inte upprepa titeln")
    return fel, varn


def hitta_kommun(namn: str, kommuner: list):
    idx = {k.casefold(): k for k in kommuner}
    n = re.sub(r"\s+(kommun|stad|municipality)$", "", namn.casefold().strip())
    if n in idx:
        return idx[n], None
    if n.endswith("s") and n[:-1] in idx:
        return idx[n[:-1]], None
    f = difflib.get_close_matches(n, list(idx), 1, 0.75)
    return None, (idx[f[0]] if f else None)


def validera(data, tillat_exempel: bool = False):
    c = Ctx(tillat_exempel)
    if not isinstance(data, dict):
        c.fel.append("data.json måste vara ett objekt")
        return c, None
    p = data.get("period") or {}
    try:
        typ, nr, ar = p["typ"], int(p["nr"]), int(p["ar"])
        if typ not in ("vecka", "manad"):
            raise ValueError
        start, slut = berakna_period(typ, nr, ar)
    except Exception:
        c.fel.append("period saknas eller är ogiltig – skapa data.json med init")
        return c, None
    if p.get("start") != start.isoformat() or p.get("slut") != slut.isoformat():
        c.fel.append(f"period.start/slut stämmer inte med {typ} {nr} {ar} ({start}–{slut}); "
                     "ändra inte periodfälten för hand")
    kd = las_datum(data.get("kontrolldatum"))
    if not kd:
        c.fel.append("kontrolldatum saknas eller är ogiltigt (ÅÅÅÅ-MM-DD)")
        kd = idag_sthlm()
    m = {"period": {"typ": typ, "nr": nr, "ar": ar, "start": start, "slut": slut},
         "kd": kd, "val": [v for v in map(txt, data.get("val") or []) if v], "regioner": []}
    for k in data:
        if not k.startswith("_") and k not in ("version", "period", "kontrolldatum", "val", "regioner"):
            c.fel.append(f"okänt toppnivåfält '{k}'")
    regioner = data.get("regioner") or []
    if not regioner:
        c.fel.append("regioner saknas")
    koder = {}
    for r in regioner:
        reg = validera_region(r, m, c)
        if reg is None:
            continue
        if reg["kod"] in koder:
            c.fel.append(f"ID-prefixet {reg['kod']} används av både {koder[reg['kod']]} och {reg['kort']}")
        koder[reg["kod"]] = reg["kort"]
        m["regioner"].append(reg)
    if len({r["kort"] for r in m["regioner"]}) != len(m["regioner"]):
        c.fel.append("två regioner har samma kortnamn")
    return c, m


def validera_region(r, m, c):
    if not isinstance(r, dict) or not txt(r.get("namn")):
        c.fel.append("en region saknar namn")
        return None
    namn = txt(r["namn"])
    kort = txt(r.get("kort")) or namn
    kod = (txt(r.get("kod")) or "").upper()
    if not re.fullmatch(r"[A-ZÅÄÖ]{1,3}", kod):
        c.fel.append(f"{kort}: kod måste vara 1–3 versaler (ID-prefix)")
    kommuner = [k for k in map(txt, r.get("kommuner") or []) if k]
    if not kommuner:
        c.varn.append(f"{kort}: kommunlistan är tom – kommun och täckning kan inte kontrolleras")
    tillatna = {"namn", "kort", "kod", "kommuner", "sokfragor", "tidigare_kallor", "kallor",
                "researchluckor", "evenemang"}
    for k in r:
        if not k.startswith("_") and k not in tillatna:
            c.fel.append(f"{kort}: okänt fält '{k}' i regionen")
    reg = {"namn": namn, "kort": kort, "kod": kod, "kommuner": kommuner,
           "sokfragor": [s for s in map(txt, r.get("sokfragor") or []) if s],
           "tidigare": [t for t in (r.get("tidigare_kallor") or []) if isinstance(t, dict)],
           "researchluckor": [s for s in map(txt, r.get("researchluckor") or []) if s],
           "kallor": [], "A": [], "B": []}
    for j, k in enumerate(r.get("kallor") or []):
        kk = validera_kalla(k, j, reg, c)
        if kk:
            reg["kallor"].append(kk)
    for i, e in enumerate(r.get("evenemang") or []):
        ev = validera_evenemang(e, i, reg, m, c)
        if ev:
            (reg["A"] if ev["status"] in A_STATUS else reg["B"]).append(ev)
    sortera_och_numrera(reg)
    kontrollera_kvot(reg, c)
    kontrollera_dubbletter(reg, c)
    berakna_tackning(reg, c)
    return reg


def validera_kalla(k, j, reg, c):
    lbl = f"{reg['kort']} källa #{j + 1}"
    if not isinstance(k, dict):
        c.fel.append(f"{lbl}: måste vara ett objekt")
        return None
    for f, v in k.items():
        if f.startswith("_"):
            continue
        if f not in MALL_KALLA:
            c.fel.append(f"{lbl}: okänt fält '{f}'")
        elif isinstance(v, str) and v.strip() == MALL_KALLA[f]:
            c.fel.append(f"{lbl}: platshållaren i '{f}' är kvar – skriv värdet eller null")
    typ, res = txt(k.get("typ")), txt(k.get("resultat"))
    if typ not in KALLTYPER:
        c.fel.append(f"{lbl}: typ måste vara en av {', '.join(KALLTYPER)}")
    if res not in RESULTAT:
        c.fel.append(f"{lbl}: resultat måste vara en av {', '.join(RESULTAT)}")
    url, fraga = txt(k.get("url")), txt(k.get("fraga"))
    if typ == "sokning" and not fraga:
        c.fel.append(f"{lbl}: sokning kräver fraga")
    if typ != "sokning" and not url:
        c.fel.append(f"{lbl}: url saknas")
    kommun = txt(k.get("kommun"))
    if kommun and reg["kommuner"]:
        hit, forslag = hitta_kommun(kommun, reg["kommuner"])
        if not hit:
            c.fel.append(f"{lbl}: kommunen '{kommun}' finns inte i {reg['namn']}"
                         + (f" – menade du {forslag}?" if forslag else ""))
        kommun = hit or kommun
    return {"typ": typ, "resultat": res, "url": url, "fraga": fraga, "kommun": kommun}


def validera_evenemang(e, i, reg, m, c):
    if not isinstance(e, dict):
        c.fel.append(f"{reg['kort']} evenemang #{i + 1}: måste vara ett objekt")
        return None
    titel = txt(e.get("titel"))
    lbl = f"{reg['kort']} evenemang #{i + 1} «{(titel or '?')[:60]}»"

    def F(s):
        c.fel.append(f"{lbl}: {s}")

    def W(s):
        c.varn.append(f"{lbl}: {s}")

    for k, v in e.items():
        if k.startswith("_"):
            continue
        if k not in MALL_EVENEMANG:
            f = difflib.get_close_matches(ascii_vikning(k), list(MALL_EVENEMANG), 1, 0.6)
            F(f"okänt fält '{k}'" + (f" – menade du '{f[0]}'?" if f else ""))
        elif isinstance(v, str) and v.strip() == MALL_EVENEMANG[k]:
            F(f"platshållaren i '{k}' är kvar – skriv värdet eller null")
    status = (txt(e.get("status")) or "").casefold()
    if status not in A_STATUS + B_STATUS:
        F(f"status måste vara en av {', '.join(A_STATUS + B_STATUS)}")
        return None
    A = status in A_STATUS
    if not titel:
        F("titel saknas")
    elif har_emoji(titel):
        W("titeln innehåller emoji eller symboler – ta bort dem")
    per = m["period"]
    ev = {"status": status, "titel": titel or "?", "nr": i + 1}

    kd_e = las_datum(e.get("kontrolldatum"))
    if kd_e is False:
        F("kontrolldatum ska vara ÅÅÅÅ-MM-DD")
    ev["kd"] = kd_e or m["kd"]
    upd = las_datum(e.get("uppdaterad"))
    if upd is False:
        F("uppdaterad ska vara ÅÅÅÅ-MM-DD")
        upd = None
    ev["uppdaterad"] = upd
    for f in ("arrangor", "malgrupp", "deltagarkrav", "pris", "anmalan_text", "kalla",
              "osakerheter", "skal", "ort", "lokal", "adress", "digital_plattform",
              "regional_koppling", "belagg", "beskrivning"):
        ev[f] = txt(e.get(f))
    if not ev["kalla"]:
        F("kalla saknas (var kandidaten hittades och verifierades)")

    d0, d1 = las_datum(e.get("datum")), las_datum(e.get("datum_slut"))
    if d0 is False:
        F("datum ska vara ÅÅÅÅ-MM-DD")
        d0 = None
    if d1 is False:
        F("datum_slut ska vara ÅÅÅÅ-MM-DD")
        d1 = None
    if d0 and d1 and d1 < d0:
        F("datum_slut ligger före datum")
        d1 = None
    if d1 == d0:
        d1 = None
    ts, te = las_tid(e.get("tid_start")), las_tid(e.get("tid_slut"))
    if ts is False:
        F("tid_start ska vara HH:MM")
        ts = None
    if te is False:
        F("tid_slut ska vara HH:MM")
        te = None
    if ts and te and not d1 and te <= ts:
        F("tid_slut måste vara senare än tid_start")
    ev.update(datum=d0, datum_slut=d1, tid_start=ts, tid_slut=te)

    lank = txt(e.get("lank"))
    if lank:
        if not re.match(r"https?://[^\s/]+\.[^\s/]+", lank):
            F("lank måste vara en fullständig http(s)-adress")
        elif urlsplit(lank).hostname and urlsplit(lank).hostname.endswith(".invalid") \
                and not c.tillat_exempel:
            F("lank är en exempellänk (.invalid)")
    ev["lank"] = lank
    ev["samlingssida"] = las_bool(e.get("samlingssida", False))
    ev["prelim"] = las_bool(e.get("preliminar_lokal", False))

    fmt = (txt(e.get("format")) or "plats").casefold()
    if fmt not in FORMAT:
        F(f"format måste vara en av {', '.join(FORMAT)}")
        fmt = "plats"
    ev["format"] = fmt
    kommun = txt(e.get("kommun"))
    if kommun and reg["kommuner"]:
        hit, forslag = hitta_kommun(kommun, reg["kommuner"])
        if hit:
            kommun = hit
        elif A and fmt != "digitalt":
            F(f"kommunen '{kommun}' ingår inte i {reg['namn']} (R1?)"
              + (f" – menade du {forslag}?" if forslag else ""))
    ev["kommun"] = kommun

    anm = (txt(e.get("anmalan")) or "").casefold()
    tf = las_datum(e.get("tidsfrist"))
    if tf is False:
        F("tidsfrist ska vara ÅÅÅÅ-MM-DD")
        tf = None
    ev["tidsfrist"] = tf
    if anm and anm not in ANMALAN:
        F(f"anmalan måste vara en av {', '.join(ANMALAN)}")
        anm = ""
    if tf and anm != "kravs":
        F("tidsfrist är satt – då måste anmalan vara \"kravs\"")
    ev["anmalan"] = anm or None

    if status == "utesluten":
        regel = (txt(e.get("regel")) or "").upper()
        if regel not in REGLER:
            F("utesluten kräver regel R1–R6 eller R9")
        ev["regel"] = regel
        if not ev["skal"]:
            F("utesluten kräver skal")
    elif txt(e.get("regel")):
        F("regel anges bara när status är utesluten")
    if status in ("obekraftad", "installt") and not ev["skal"]:
        F(f"{status} kräver skal")
    if not A:
        ev.update(typ=None, villkor=[], pris_tagg=txt(e.get("pris_tagg")), dropin=None, ovrigt=None,
                  kategori=txt(e.get("kategori")))
        return ev

    # ---- endast kalenderevenemang (verifierad / vantelista)
    typ = txt(e.get("typ"))
    kanon = {t.casefold(): t for t in TYPER}
    if not typ or typ.casefold() not in kanon:
        F(f"typ måste vara en av: {', '.join(TYPER)}")
    ev["typ"] = kanon.get((typ or "").casefold(), typ or "?")
    kat = txt(e.get("kategori"))
    if kat not in KATEGORIER:
        F(f"kategori måste vara en av: {', '.join(KATEGORIER)}")
    ev["kategori"] = kat
    if not ev["arrangor"]:
        W("arrangor saknas – arrangörsraden utelämnas på kortet")
    if not ev["beskrivning"]:
        F("beskrivning saknas")
    else:
        bf, bv = kontrollera_beskrivning(ev["beskrivning"], ev["titel"])
        for s in bf:
            F(s)
        for s in bv:
            W(s)
    if not d0:
        F("datum saknas")
    elif d0 > per["slut"] or (d1 or d0) < per["start"]:
        F(f"datum {d0}{('–' + str(d1)) if d1 else ''} ligger utanför perioden "
          f"{per['start']}–{per['slut']} (R1)")
    if not ts:
        W("tid_start saknas – kortet visar 'Tid ej angiven'; kontrollera att arrangören inte anger någon tid")
    if fmt in ("plats", "hybrid"):
        if not ev["ort"]:
            F("ort saknas")
        if not kommun and reg["kommuner"]:
            F("kommun saknas")
        if ev["adress"] and re.search(r"\b\d{3} ?\d{2}\b", ev["adress"]):
            F("adress ska anges utan postnummer")
    if fmt == "digitalt" and not ev["regional_koppling"] and ev["typ"] not in DIGITALT_UTAN_KOPPLING:
        F("digitalt evenemang kräver regional_koppling (annars R6), utom jobbmässor och rekryteringsträffar")
    if not anm:
        F("anmalan saknas (kravs, ingen eller okand)")
    if tf and d0 and tf > (d1 or d0):
        W("tidsfristen ligger efter evenemanget – kontrollera")
    if not lank:
        F("lank saknas")
    if not ev["belagg"] or len(ev["belagg"]) < 15:
        F("belagg saknas – citera arrangörens sida ordagrant (datum, år och starttid)")
    elif d0:
        b, bl = ev["belagg"], ev["belagg"].casefold()
        if str(d0.day) not in b or (ts and str(int(ts[:2])) not in b):
            W("belagg verkar inte innehålla datum och starttid – kontrollera")
        # Året får räknas fram om sidan är uppdaterad under periodens år eller halvåret före.
        ar_framraknat = str(d0.year) not in b
        if ar_framraknat:
            if not upd:
                F("året står inte i belagg – ange uppdaterad (sidans publicerings- eller "
                  "uppdateringsdatum), annars är evenemanget obekraftad")
            elif upd > ev["kd"] or (upd.year not in (per["start"].year, per["slut"].year)
                                    and upd < per["start"] - dt.timedelta(days=183)):
                F(f"sidan uppdaterades {upd} – för gammal för att året ska få räknas fram (obekraftad)")
        # En angiven veckodag måste stämma med datumet; annars är sidan troligen gammal.
        dagar = {VECKODAGAR[(d0 + dt.timedelta(n)).weekday()] for n in range(min(((d1 or d0) - d0).days + 1, 7))}
        if any(v in bl for v in VECKODAGAR) and not any(v in bl for v in dagar):
            (F if ar_framraknat else W)("veckodagen i belagg stämmer inte med datumet – gammal sida?")

    pt = txt(e.get("pris_tagg"))
    if pt:
        pt = tusental(pt)
        if not any(r.fullmatch(pt) for r in PRIS_RE):
            F(f"pris_tagg '{pt}' följer inte vokabulären (Gratis | Fri entré | 150 kr | "
              "150 kr medlem / 250 kr övriga ex moms)")
        elif pt == "Fri entré" and ev["typ"] != "Mässa":
            F("Fri entré används bara för mässor – skriv Gratis")
    ev["pris_tagg"] = pt

    vk = e.get("villkor") or []
    if isinstance(vk, str):
        vk = [vk]
    villkor = []
    for v in map(txt, vk):
        if not v:
            continue
        v = tusental(v)
        if v.casefold().startswith("endast medlemmar"):
            F("'Endast medlemmar' går inte: slutna medlemsevenemang utesluts (R4)")
            continue
        hit = [n for n, r in enumerate(VILLKOR_RE) if r.fullmatch(v)]
        if not hit:
            F(f"villkor '{v}' följer inte vokabulären")
            continue
        if v == "Väntelista" and status != "vantelista":
            F("villkoret Väntelista kräver status vantelista")
            continue
        villkor.append((hit[0], v))
    if status == "vantelista" and not any(v == "Väntelista" for _, v in villkor):
        villkor.append((4, "Väntelista"))
    ev["villkor"] = [v for _, v in sorted(villkor, key=lambda x: x[0])]

    ov = e.get("ovrigt")
    if isinstance(ov, list):
        if len(ov) > 1:
            F("ovrigt: högst en tagg")
        ov = ov[0] if ov else None
    ov = txt(ov)
    if ov and len(ov) > 40:
        F("ovrigt får vara högst 40 tecken")
    ev["ovrigt"] = ov.rstrip(".") if ov else None

    dropin = las_intervall(e.get("dropin"))
    if dropin is False:
        F("dropin ska vara HH:MM–HH:MM")
        dropin = None
    ev["dropin"] = dropin

    ak = txt(e.get("arrangor_kort")) or ev["arrangor"]
    if ak:
        ak = re.sub(r"\bi samarbete med\b", "med", ak, flags=re.I)
        if len(ak) > 90:
            F(f"arrangörsraden har {len(ak)} tecken (max 90) – skriv arrangor_kort")
    ev["arr_kort"] = ak
    harled_kort(ev, per)
    return ev


def harled_kort(ev: dict, per: dict):
    """Fält som skriptet bestämmer: visningsdag, tid, plats och taggar."""
    d0, d1 = ev["datum"], ev["datum_slut"]
    ev["visning"] = max(d0, per["start"]) if d0 else None
    ev["tid_text"] = (ev["tid_start"] + (f"–{ev['tid_slut']}" if ev["tid_slut"] else "")) \
        if ev["tid_start"] else "Tid ej angiven"
    if ev["format"] == "digitalt":
        ev["plats_text"] = "Digitalt"
    else:
        delar = []
        for d in (ev["lokal"], ev["adress"], ev["ort"]):
            if d and d.casefold() not in (x.casefold() for x in delar):
                delar.append(d)
        ev["plats_text"] = ", ".join(delar) + (" (preliminärt)" if ev["prelim"] else "")
    taggar = [(ev["typ"], False)]
    if ev["format"] == "digitalt":
        taggar.append(("Digitalt", False))
    elif ev["format"] == "hybrid":
        taggar.append(("Även digitalt", False))
    if ev["dropin"]:
        taggar.append((f"Drop-in {ev['dropin']}", False))
    if d0 and d1:
        taggar.append((f"Pågår {intervall(d0, d1, med_ar=False)}", False))
    if ev["pris_tagg"]:
        taggar.append((ev["pris_tagg"], False))
    taggar += [(v, False) for v in ev["villkor"]]
    if ev["ovrigt"]:
        taggar.append((ev["ovrigt"], False))
    if ev["anmalan"] == "kravs":
        tf = ev["tidsfrist"]
        taggar.append((f"Anmäl senast {datum_text(tf)}" if tf and tf >= ev["kd"]
                       else "Anmälan krävs", True))
    ev["taggar"] = taggar


def sortera_och_numrera(reg: dict):
    reg["A"].sort(key=lambda e: (e["visning"] or dt.date.max, e["tid_start"] or "99:99",
                                 sv_nyckel(e["titel"])))
    for n, e in enumerate(reg["A"], 1):
        e["id"] = f"{reg['kod']}{n}"
    reg["B"].sort(key=lambda e: (e["datum"] or dt.date.max, e["tid_start"] or "99:99",
                                 sv_nyckel(e["titel"])))
    for n, e in enumerate(reg["B"], 1):
        e["id"] = f"{reg['kod']}U{n}"


def kontrollera_dubbletter(reg: dict, c: Ctx):
    sedda, lankar = {}, {}
    alla = sorted(reg["A"] + reg["B"], key=lambda e: e["nr"])
    par = lambda a, b: f"evenemang #{a['nr']} och #{b['nr']} «{b['titel']}»"
    for e in alla:
        if not e["datum"]:
            continue
        nyckel = (e["datum"], e["tid_start"], norm_jmf(e["titel"]))
        dubbel = nyckel in sedda
        if dubbel:
            c.fel.append(f"{reg['kort']}: dubblett – {par(sedda[nyckel], e)} har samma titel, datum och tid. "
                         "Behåll en.")
        sedda.setdefault(nyckel, e)
        if e["lank"]:
            lk = (e["datum"], e["tid_start"], e["lank"].rstrip("/"))
            if lk in lankar and not dubbel:
                c.fel.append(f"{reg['kort']}: dubblett – {par(lankar[lk], e)} har samma länk, datum och tid. "
                             "Behåll en.")
            lankar.setdefault(lk, e)
    for i, a in enumerate(alla):
        for b in alla[i + 1:]:
            if a["datum"] and a["datum"] == b["datum"] and norm_jmf(a["titel"]) != norm_jmf(b["titel"]) \
                    and difflib.SequenceMatcher(None, norm_jmf(a["titel"]), norm_jmf(b["titel"])).ratio() > 0.85:
                c.varn.append(f"{reg['kort']}: möjlig dubblett – #{a['nr']} «{a['titel']}» och "
                              f"#{b['nr']} «{b['titel']}» samma dag")


def kontrollera_kvot(reg: dict, c: Ctx):
    """Kompetensutveckling + Företagande får vara högst 40 % (högst 2 per 3 arbetsmarknadskort)."""
    arb = [e for e in reg["A"] if e.get("kategori") == "Arbetsmarknad"]
    ovr = [e for e in reg["A"] if e.get("kategori") in ("Kompetensutveckling", "Företagande")]
    tak = 2 * len(arb) // 3
    if len(ovr) > tak:
        ordning = sorted(ovr, key=lambda e: (e["kategori"] != "Företagande", -(e["visning"] or dt.date.min).toordinal(),
                                              e["tid_start"] or ""), reverse=False)
        stryk = ordning[:len(ovr) - tak]
        c.fel.append(f"{reg['kort']}: 40 %-kvoten överskrids ({len(ovr)} kompetensutveckling/företagande, tak {tak} "
                     f"med {len(arb)} arbetsmarknadskort). Sätt status utesluten, regel R9 på: "
                     + "; ".join(f"#{e['nr']} «{e['titel']}»" for e in stryk))


def berakna_tackning(reg: dict, c: Ctx):
    kallor, kommuner = reg["kallor"], reg["kommuner"]
    per_kommun = {k: set() for k in kommuner}      # ej läsbara räknas som kontrollerade men listas
    for k in kallor:
        if k["kommun"] in per_kommun:
            per_kommun[k["kommun"]].add(k["typ"])
    t = {
        "kommun_saknas": [k for k in kommuner if "kommun" not in per_kommun[k]],
        "komvux_saknas": [k for k in kommuner if "komvux" not in per_kommun[k]],
        "typer_saknas": [n for n in REGIONALA_KALLTYPER if n not in {k["typ"] for k in kallor}],
        "ej_lasbara": [k["url"] or k["fraga"] for k in kallor if k["resultat"] == "ej_lasbar"],
    }
    gjorda = {norm_jmf(k["fraga"]) for k in kallor if k["typ"] == "sokning" and k["fraga"]}
    t["sok_saknas"] = [s for s in reg["sokfragor"] if norm_jmf(s) not in gjorda]
    nu = {(k["url"] or "").rstrip("/") for k in kallor}
    t["tidigare_saknas"] = [u for u in ((txt(x.get("url")) or "").rstrip("/") for x in reg["tidigare"])
                            if u and u not in nu]
    med_traff = {e["kommun"] for e in reg["A"] + reg["B"] if e["kommun"]}
    t["utan_traff"] = [k for k in kommuner if k not in med_traff]
    reg["tackning"] = t

    luckor = list(reg["researchluckor"])
    if not reg["A"]:
        luckor.append("inga verifierade evenemang – ingen kalender skapad")
    if t["ej_lasbara"]:
        luckor.append("gick inte att läsa: " + ", ".join(t["ej_lasbara"]))
    if t["kommun_saknas"]:
        luckor.append("kommunens evenemangs-/näringslivssida ej kontrollerad: " + ", ".join(t["kommun_saknas"]))
    if t["komvux_saknas"]:
        luckor.append("komvux ej kontrollerat: " + ", ".join(t["komvux_saknas"]))
    if t["typer_saknas"]:
        luckor.append("källtyp ej kontrollerad: " + ", ".join(REGIONALA_KALLTYPER[n] for n in t["typer_saknas"]))
    if t["sok_saknas"]:
        luckor.append(f"standardsökningar ej körda: {len(t['sok_saknas'])} av {len(reg['sokfragor'])}")
    if t["tidigare_saknas"]:
        luckor.append(f"källor från förra körningen ej kontrollerade: {len(t['tidigare_saknas'])}")
    if t["utan_traff"] and kommuner:
        luckor.append("kommuner utan träffar: " + ", ".join(t["utan_traff"]))
    reg["luckor"] = luckor
    for s in ("kommun_saknas", "komvux_saknas", "typer_saknas", "sok_saknas", "tidigare_saknas"):
        if t[s]:
            c.varn.append(f"{reg['kort']}: täckning ofullständig ({s.replace('_', ' ')}: {len(t[s])}) "
                          "– kontrollera källorna eller låt det stå som researchlucka")


def visa_meddelanden(c: Ctx):
    for s in c.fel:
        print(f"FEL    {s}")
    for s in c.varn:
        print(f"VARNING {s}")


def visa_tackning(m: dict):
    for reg in m["regioner"]:
        t, n = reg["tackning"], len(reg["kommuner"])
        print(f"\nTÄCKNING {reg['kort']} ({n} kommuner, {len(reg['kallor'])} loggade källor)")
        if n:
            print(f"  kommunsida  {n - len(t['kommun_saknas'])}/{n}"
                  + (f"  saknas: {', '.join(t['kommun_saknas'])}" if t["kommun_saknas"] else ""))
            print(f"  komvux      {n - len(t['komvux_saknas'])}/{n}"
                  + (f"  saknas: {', '.join(t['komvux_saknas'])}" if t["komvux_saknas"] else ""))
        print("  regionalt   " + "  ".join(f"{k} {'✓' if k not in t['typer_saknas'] else '✗'}"
                                         for k in REGIONALA_KALLTYPER))
        ns = len(reg["sokfragor"])
        print(f"  sökningar   {ns - len(t['sok_saknas'])}/{ns}")
        if reg["tidigare"]:
            print(f"  förra körningens källor ej kontrollerade: {len(t['tidigare_saknas'])}")
        if t["ej_lasbara"]:
            print(f"  ej läsbara  {len(t['ej_lasbara'])}")
        ob = sum(e["status"] == "obekraftad" for e in reg["B"])
        print(f"  evenemang   A {len(reg['A'])} (kalender) · B {len(reg['B'])} "
              f"(obekräftade {ob}, uteslutna/inställda {len(reg['B']) - ob})")

# ---------------------------------------------------------------- Excel


KOLUMNER = [("Internt ID", 9), ("Namn", 38), ("Kategori", 18), ("Arrangör", 30), ("Kort beskrivning", 48),
            ("Datum", 11), ("Tider", 16), ("Ort", 14), ("Lokal", 24),
            ("Adress eller digitalt format", 26), ("Målgrupp", 24), ("Deltagarkrav", 20),
            ("Pris", 20), ("Anmälningskrav", 20), ("Tidsfrist", 18), ("Status", 30),
            ("Verifierad direktlänk", 44), ("Källa", 30), ("Kontrolldatum", 13), ("Osäkerheter", 44)]


def excel_rad(e: dict) -> list:
    A = e["status"] in A_STATUS
    d = e["visning"] if A else e["datum"]
    if e["tid_start"]:
        tider = e["tid_start"] + (f"–{e['tid_slut']}" if e["tid_slut"] else "")
        if e["datum_slut"]:
            tider += f" ({intervall(e['datum'], e['datum_slut'], med_ar=False)})"
    else:
        tider = EJ
    if e["format"] == "digitalt":
        ort = lokal = "Digitalt"
        adress = f"Digitalt ({e['digital_plattform']})" if e["digital_plattform"] else "Digitalt"
    else:
        ort = e["ort"] or EJ
        lokal = (e["lokal"] + (" (preliminärt)" if e["prelim"] else "")) if e["lokal"] else EJ
        adress = e["adress"] or EJ
        if e["format"] == "hybrid":
            adress += " + digitalt" + (f" ({e['digital_plattform']})" if e["digital_plattform"] else "")
    anm = e["anmalan_text"] or {"kravs": "Ja", "ingen": "Nej"}.get(e["anmalan"], EJ)
    tf = e["tidsfrist"]
    tidsfrist = (tf.isoformat() + (" (passerad)" if tf < e["kd"] else "")) if tf else EJ
    status = {"verifierad": f"Planerat (verifierat {e['kd'].isoformat()})",
              "vantelista": "Fullbokat – väntelista", "installt": "Inställt",
              "obekraftad": "Obekräftad"}.get(e["status"])
    if e["status"] == "utesluten":
        status = f"Utesluten: {e.get('regel', '')} – {e['skal']}"
    osak = []
    if e["status"] in ("obekraftad", "installt") and e["skal"]:
        osak.append(e["skal"])
    if A:
        if e["samlingssida"]:
            osak.append("Länken går till en samlingssida")
        if e["prelim"]:
            osak.append("Preliminär lokal")
        if tf and tf < e["kd"]:
            osak.append("Anmälningsfristen har passerat (kortet visar 'Anmälan krävs')")
        if not e["pris_tagg"] and not e["pris"]:
            osak.append("Pris ej angivet")
    if e["osakerheter"]:
        osak.append(e["osakerheter"])
    return [e["id"], e["titel"], e.get("kategori") or EJ, e["arrangor"] or EJ, e["beskrivning"] or EJ, d or EJ, tider,
            ort, lokal, adress, e["malgrupp"] or EJ, e["deltagarkrav"] or EJ,
            e["pris"] or e["pris_tagg"] or EJ, anm, tidsfrist, status, e["lank"] or EJ,
            e["kalla"] or EJ, e["kd"], "; ".join(osak) if osak else "Inga"]


def skriv_excel(path: Path, m: dict, et: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    F = Font(name="Arial", size=10)
    FB = Font(name="Arial", size=10, bold=True)
    FL = Font(name="Arial", size=10, color="0563C1", underline="single")
    fyll = PatternFill("solid", fgColor="DCE6F0")
    radbryt = Alignment(wrap_text=True, vertical="top")
    p = m["period"]
    iv = intervall(p["start"], p["slut"], med_ar=False)
    for reg in m["regioner"]:
        ws = wb.create_sheet(title=re.sub(r"[\[\]:*?/\\]", "", reg["kort"])[:31])

        def skriv(rad, kol, varde, font=F, align=None):
            cell = ws.cell(row=rad, column=kol, value=varde)
            cell.font = font
            if align:
                cell.alignment = align
            return cell

        skriv(1, 1, f"Källunderlag – {reg['namn']} – {et['namn']}, {p['ar']} ({iv})", FB)
        skriv(2, 1, f"Kontrolldatum {m['kd'].isoformat()}. 'ej angivet' = uppgiften saknas hos källan. "
                    "Endast avsnitt A ingår i kalendern.")
        rad = 4
        for rubrik, lista, tom in (
                (f"A. VERIFIERADE EVENEMANG – ingår i kalendern ({len(reg['A'])} st)", reg["A"],
                 "Inga verifierade evenemang hittades för perioden."),
                (f"B. OBEKRÄFTADE KANDIDATER OCH UTESLUTNA – ingår EJ i kalendern ({len(reg['B'])} st)",
                 reg["B"], "Inga.")):
            skriv(rad, 1, rubrik, FB)
            rad += 1
            for k, (namn, _) in enumerate(KOLUMNER, 1):
                skriv(rad, k, namn, FB, radbryt).fill = fyll
            rad += 1
            if not lista:
                skriv(rad, 1, tom)
                rad += 1
            for e in lista:
                for k, v in enumerate(excel_rad(e), 1):
                    if isinstance(v, str):
                        v = v.replace("­", "")        # mjuka bindestreck hör bara till kortet
                    cell = skriv(rad, k, v, F, radbryt)
                    if isinstance(v, dt.date):
                        cell.number_format = "yyyy-mm-dd"
                    if KOLUMNER[k - 1][0] == "Verifierad direktlänk" and e["lank"]:
                        cell.hyperlink = e["lank"]
                        cell.font = FL
                rad += 1
            rad += 1
        skriv(rad, 1, "Researchluckor: " + ("; ".join(reg["luckor"]) if reg["luckor"] else "inga"))
        for k, (_, bredd) in enumerate(KOLUMNER, 1):
            ws.column_dimensions[get_column_letter(k)].width = bredd
    wb.properties.creator = f"bygg_kalender.py {VERSION}"
    wb.save(path)

# ---------------------------------------------------------------- HTML


def kort_html(e: dict) -> str:
    org = f'<div class="o">{esc(e["arr_kort"])}</div>' if e["arr_kort"] else ""
    taggar = "".join(f'<span class="tag{" dl" if dl else ""}">{esc(t)}</span>' for t, dl in e["taggar"])
    return (f'<a class="card" href="{esc(e["lank"])}" data-id="{e["id"]}">\n'
            f'  <div class="top">\n'
            f'    <div class="l"><div class="t">{esc(e["titel"])}</div>{org}</div>\n'
            f'    <div class="r">\n'
            f'      <div class="tm">{SVG_TID}<span>{esc(e["tid_text"])}</span></div>\n'
            f'      <div class="pl">{SVG_PLATS}<span>{esc(e["plats_text"])}</span></div>\n'
            f'    </div>\n'
            f'  </div>\n'
            f'  <p class="d">{esc(e["beskrivning"])}</p>\n'
            f'  <div class="tags">{taggar}</div>\n'
            f'</a>')


def indrag(s: str, n: int) -> str:
    return "\n".join(" " * n + rad for rad in s.split("\n"))


def html_dokument(titel: str, kropp: str, skript: bool = True) -> str:
    return ('<!doctype html><html lang="sv"><head><meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
            f'<title>{esc(titel)}</title>\n<style>\n{CSS}\n</style></head><body>\n\n'
            f'{kropp}\n\n' + (ZOOM_JS + '\n' if skript else '') + '</body></html>\n')


def rubrikblock(reg: dict, m: dict, et: dict) -> str:
    p = m["period"]
    return (f'  <h1>{esc(et["h1"])}</h1>\n'
            f'  <div class="sub"><b>{esc(reg["namn"].upper())}</b> · '
            f'{esc(intervall(p["start"], p["slut"]))}</div>')


def dagblock(cls: str, dag: dt.date, kort: list) -> str:
    delar = [f'  <div class="{cls}">', f'    <div class="dh"><span>{esc(dagrubrik(dag))}</span></div>']
    delar += [indrag(kort_html(e), 4) for e in kort]
    delar.append('  </div>')
    return "\n".join(delar)


def gruppera_dagar(A: list) -> list:
    dagar = []
    for e in A:
        if not dagar or dagar[-1][0] != e["visning"]:
            dagar.append((e["visning"], []))
        dagar[-1][1].append(e)
    return dagar


MAT_JS = """() => {
  const mm = 25.4 / 96, h = el => el.getBoundingClientRect().height * mm;
  const s = document.querySelector('section.sheet');
  return {h1: h(s.querySelector('h1')), sub: h(s.querySelector('.sub')),
          dagar: [...s.querySelectorAll('.day')].map(d => ({dh: h(d.querySelector('.dh')),
                   kort: [...d.querySelectorAll('a.card')].map(h)}))};
}"""


def packa(matt: dict, dagar: list):
    """v2:s paginering: dagsblock i tur och ordning, rubrik bara om minst ett kort ryms."""
    tillganglig = SIDHOJD - SAKERHET
    sidor, sida = [], []
    y = matt["h1"] + MARG_RUBRIK + matt["sub"] + MARG_UNDERRAD
    botten = []
    for di, (_, kort) in enumerate(dagar):
        dh, hojder = matt["dagar"][di]["dh"], matt["dagar"][di]["kort"]
        ci = 0
        while ci < len(kort):
            glapp = MARG_DAG if sida else 0.0
            if y + glapp + dh + MARG_DH + hojder[ci] > tillganglig:
                if not sida:
                    fel_avbryt(f"kortet {kort[ci]['id']} är högre än en sida")
                sidor.append(sida)
                botten.append(y)
                sida, y = [], 0.0
                continue
            y += glapp + dh + MARG_DH + hojder[ci]
            block = [ci]
            ci += 1
            while ci < len(kort) and y + MARG_KORT + hojder[ci] <= tillganglig:
                y += MARG_KORT + hojder[ci]
                block.append(ci)
                ci += 1
            sida.append((di, block))
            if ci < len(kort):
                sidor.append(sida)
                botten.append(y)
                sida, y = [], 0.0
    sidor.append(sida)
    botten.append(y)
    return sidor, botten


QA_JS = """() => {
  const MM = 96 / 25.4, LH = %(lh)f;
  const sheets = [...document.querySelectorAll('section.sheet')].map(s => {
    const sr = s.getBoundingClientRect();
    const delar = [...s.querySelectorAll('h1, .sub, .dh, a.card')];
    const sista = delar[delar.length - 1];
    return {hojd: sr.height / MM, marginal: sista ? (sr.bottom - sista.getBoundingClientRect().bottom) / MM : null,
            sistKort: sista ? sista.matches('a.card') : false, overflow: s.scrollHeight - s.clientHeight,
            kort: [...s.querySelectorAll('a.card')].map(a => a.dataset.id)};
  });
  const h1 = document.querySelector('h1');
  const kort = [...document.querySelectorAll('a.card')].map(a => {
    const ar = a.getBoundingClientRect(), d = a.querySelector('.d'), tm = a.querySelector('.tm');
    const over = [...a.querySelectorAll('.t, .o, .r, .tm, .pl, .tags, .tag')]
      .filter(el => el.scrollWidth > el.clientWidth + 1).map(el => el.className);
    const utanfor = [...a.querySelectorAll('*')].some(el => {
      const r = el.getBoundingClientRect();
      return r.right > ar.right + 0.5 || r.left < ar.left - 0.5 || r.bottom > ar.bottom + 0.5; });
    const kol = a.querySelector('.l').getBoundingClientRect().width / a.querySelector('.r').getBoundingClientRect().width;
    return {id: a.dataset.id, href: a.href, rader: d.getBoundingClientRect().height / LH,
            tidBryts: tm.scrollWidth > tm.clientWidth, over, utanfor, kol, bredd: ar.width / MM};
  });
  return {sheets, h1: {sw: h1.scrollWidth, cw: h1.clientWidth, h: h1.getBoundingClientRect().height}, kort};
}""" % {"lh": RAD_D}


def vanta_typsnitt(page):
    page.evaluate("document.fonts.ready.then(() => true)")


def typsnitt_info(page) -> dict:
    ut = {}
    try:
        cdp = page.context.new_cdp_session(page)
        cdp.send("DOM.enable")
        cdp.send("CSS.enable")
        rot = cdp.send("DOM.getDocument", {"depth": 0})["root"]["nodeId"]
        for sel in ("h1", ".sub b", ".t", ".d", ".tag"):
            nid = cdp.send("DOM.querySelector", {"nodeId": rot, "selector": sel})["nodeId"]
            if nid:
                ut[sel] = [(f["familyName"], f.get("postScriptName", ""))
                           for f in cdp.send("CSS.getPlatformFontsForNode", {"nodeId": nid})["fonts"]]
        cdp.detach()
    except Exception as ex:          # CDP saknas (t.ex. annan webbläsare)
        ut["fel"] = str(ex)
    return ut


def starta_chromium(pw):
    try:
        return pw.chromium.launch()
    except Exception as ex:
        for kandidat in (os.environ.get("CHROMIUM_PATH"), "/opt/pw-browsers/chromium",
                         shutil.which("chromium"), shutil.which("chromium-browser"),
                         shutil.which("google-chrome")):
            if kandidat and Path(kandidat).exists():
                try:
                    return pw.chromium.launch(executable_path=kandidat)
                except Exception:
                    pass
        fel_avbryt(f"Chromium kunde inte startas ({str(ex).splitlines()[0]}). "
                   "Kör: python -m playwright install chromium  (eller sätt CHROMIUM_PATH)")

# ---------------------------------------------------------------- PDF och raster


def pdf_kontroll(path: Path) -> dict:
    from pypdf import PdfReader
    r = PdfReader(str(path))
    ut = {"sidor": len(r.pages), "a4": [], "lankar": [], "typsnitt": {}}
    sedda = set()

    def objekt(v):
        return v.get_object() if v is not None else {}

    def fonter(res):
        if not res:
            return
        res = res.get_object()
        for _, f in objekt(res.get("/Font")).items():
            f = f.get_object()
            namn = str(f.get("/BaseFont", "?")).lstrip("/")
            typ = f.get("/Subtype")
            if typ == "/Type3":
                ut["typsnitt"][namn] = True
                continue
            desk = f["/DescendantFonts"][0].get_object().get("/FontDescriptor") if typ == "/Type0" \
                else f.get("/FontDescriptor")
            desk = desk.get_object() if desk else {}
            ut["typsnitt"][namn] = any(k in desk for k in ("/FontFile", "/FontFile2", "/FontFile3"))
        for _, x in objekt(res.get("/XObject")).items():
            idnum = getattr(x, "idnum", None)
            if idnum is not None:
                if idnum in sedda:
                    continue
                sedda.add(idnum)
            x = x.get_object()
            if x.get("/Subtype") == "/Form":
                fonter(x.get("/Resources"))

    for i, sida in enumerate(r.pages):
        b, h = float(sida.mediabox.width), float(sida.mediabox.height)
        if abs(b - 595.28) > 1.5 or abs(h - 841.89) > 1.5:
            ut["a4"].append(f"sida {i + 1}: {b:.0f}×{h:.0f} pt")
        lankar = []
        for a in objekt(sida.get("/Annots")) or []:
            a = a.get_object()
            if a.get("/Subtype") != "/Link" or "/A" not in a:
                continue
            uri = a["/A"].get_object().get("/URI")
            x1, y1, x2, y2 = (float(v) for v in a["/Rect"])
            lankar.append({"uri": str(uri) if uri is not None else None,
                           "bredd": abs(x2 - x1) * 25.4 / 72, "botten": min(y1, y2) * 25.4 / 72})
        ut["lankar"].append(lankar)
        fonter(sida.get("/Resources"))
    return ut


def rastrera(pdf: Path, qa: Path, namn: str, dpi: int = 100) -> dict:
    ut = {"metod": None, "bilder": [], "horn": []}
    bilder = []
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(pdf))
        for i in range(len(doc)):
            img = doc[i].render(scale=dpi / 72).to_pil().convert("RGB")
            p = qa / f"{namn}_s{i + 1}.png"
            img.save(p)
            bilder.append((p, img))
        doc.close()
        ut["metod"] = "pypdfium2"
    except ImportError:
        if shutil.which("pdftoppm"):
            bas = qa / f"{namn}_tmp"
            subprocess.run(["pdftoppm", "-r", str(dpi), "-png", str(pdf), str(bas)], check=True)
            filer = sorted(qa.glob(f"{namn}_tmp-*.png"), key=lambda q: int(q.stem.rsplit("-", 1)[1]))
            for i, f in enumerate(filer, 1):
                p = qa / f"{namn}_s{i}.png"
                f.replace(p)
                try:
                    from PIL import Image
                    bilder.append((p, Image.open(p).convert("RGB")))
                except ImportError:
                    bilder.append((p, None))
            ut["metod"] = "pdftoppm"
    for p, img in bilder:
        ut["bilder"].append(p)
        if img is None:
            ut["horn"].append(None)
            continue
        w, h = img.size
        horn = [img.getpixel(xy) for xy in ((2, 2), (w - 3, 2), (2, h - 3), (w - 3, h - 3))]
        ut["horn"].append(all(r < 150 and g < 175 and bl < 210 and bl > r for r, g, bl in horn))
    return ut

# ---------------------------------------------------------------- bygg


def bygg_region(browser, reg: dict, m: dict, et: dict, ut: Path, qa: Path, res: dict, miljo: dict):
    fil = f"{et['prefix']}_{filnamn(reg['kort'])}_Evenemang"
    html_p, pdf_p = ut / f"{fil}.html", ut / f"{fil}.pdf"
    dagar = gruppera_dagar(reg["A"])
    p = m["period"]
    titel = f"EVENEMANG · {et['titel']} · {reg['namn']} · {intervall(p['start'], p['slut'])}"

    page = browser.new_page(viewport={"width": 1280, "height": 1000})
    # 1. mät alla delar i 190 mm bredd (print)
    matsida = ('<section class="sheet" style="height:auto;overflow:visible">\n'
               + rubrikblock(reg, m, et) + "\n"
               + "\n".join(dagblock("day", d, k) for d, k in dagar) + "\n</section>")
    page.emulate_media(media="print")
    page.set_content(html_dokument(titel, matsida, skript=False))
    vanta_typsnitt(page)
    matt = page.evaluate(MAT_JS)
    plan, botten = packa(matt, dagar)
    res["sidor"] = len(plan)

    # 2. skriv HTML: en <section class="sheet"> per planerad sida
    sektioner = []
    for si, sida in enumerate(plan):
        delar = ['<section class="sheet">']
        if si == 0:
            delar.append(rubrikblock(reg, m, et))
        for bi, (di, idx) in enumerate(sida):
            delar.append(dagblock("day first" if bi == 0 else "day", dagar[di][0],
                                  [dagar[di][1][i] for i in idx]))
        delar.append('</section>')
        sektioner.append("\n".join(delar))
    html_p.write_text(html_dokument(titel, "\n".join(sektioner)), encoding="utf-8")

    # 3. kontroller i webbläsaren (print)
    page.emulate_media(media="screen")
    page.goto(html_p.resolve().as_uri())
    page.wait_for_load_state("load")
    page.emulate_media(media="print")
    vanta_typsnitt(page)
    dom = page.evaluate(QA_JS)
    if "typsnitt" not in miljo:
        miljo["typsnitt"] = typsnitt_info(page)
    page.emulate_media(media="screen")

    # 4. PDF
    page.pdf(path=str(pdf_p), prefer_css_page_size=True, print_background=True, scale=1,
             margin={"top": "0", "right": "0", "bottom": "0", "left": "0"},
             display_header_footer=False)
    page.close()

    # 5. mobil
    mob = browser.new_page(viewport={"width": 390, "height": 844})
    mob.goto(html_p.resolve().as_uri())
    mob.wait_for_load_state("load")
    sw = mob.evaluate("document.documentElement.scrollWidth")
    mob.close()

    pk = pdf_kontroll(pdf_p)
    rk = rastrera(pdf_p, qa, fil)
    res["bilder"] = rk["bilder"]
    miljo.setdefault("raster", rk["metod"])
    miljo.setdefault("pdf_typsnitt", {}).update(pk["typsnitt"])

    K = lambda namn, ok, detalj="": res["kontroller"].append((namn, bool(ok), detalj))
    n_html, n_pdf, n_xls = len(dom["kort"]), sum(len(s) for s in pk["lankar"]), len(reg["A"])
    K("Data: antal kort i HTML = länkar i PDF = rader i Excel A", n_html == n_pdf == n_xls,
      f"{n_html} / {n_pdf} / {n_xls}")
    K("Data: datum inom perioden, inga dubbletter, unika ID", True, "validerat")
    bad = [f"{k['id']} ({round(k['rader'])} rader)" for k in dom["kort"] if round(k["rader"]) > 2]
    K("Layout: varje beskrivning högst 2 rader", not bad, ", ".join(bad))
    bad = [k["id"] for k in dom["kort"] if k["tidBryts"]]
    K("Layout: tiden bryts inte", not bad, ", ".join(bad))
    h1 = dom["h1"]
    K("Layout: huvudrubriken på en rad", h1["sw"] <= h1["cw"] and h1["h"] < 1.5 * RAD_H1,
      f"scrollWidth {h1['sw']} / clientWidth {h1['cw']}")
    bad = [f"sida {i + 1}: {s['hojd']:.2f} mm" for i, s in enumerate(dom["sheets"]) if abs(s["hojd"] - 297) > 0.1]
    K("Layout: varje .sheet är 297 mm hög", not bad, ", ".join(bad))
    marg = [s["marginal"] for s in dom["sheets"]]
    bad = [f"sida {i + 1}: {v:.1f} mm" for i, v in enumerate(marg) if v is None or v < 10.0]
    K("Layout: sista kortet slutar minst 10 mm ovanför underkanten", not bad,
      ", ".join(bad) or "minst " + f"{min(marg):.1f} mm")
    bad = [f"sida {i + 1}" for i, s in enumerate(dom["sheets"]) if not s["sistKort"]]
    K("Layout: ingen dagsrubrik står ensam sist på en sida", not bad, ", ".join(bad))
    bad = [f"sida {i + 1}" for i, s in enumerate(dom["sheets"]) if s["overflow"] > 1]
    bad += [f"{k['id']} ({', '.join(k['over']) or 'utanför kortet'})" for k in dom["kort"]
            if k["over"] or k["utanfor"]]
    K("Layout: inget innehåll överlappar eller är dolt", not bad, ", ".join(bad))
    bad = [f"{k['id']} ({k['kol']:.2f}:1)" for k in dom["kort"] if abs(k["kol"] - 2) > 0.02]
    K("Layout: kolumnerna i kortets överdel har proportionen 2:1", not bad,
      ", ".join(bad) + (" – ett för långt ord i titel eller plats; korta eller avstava" if bad else ""))
    K("Layout: sidindelning enligt plan", [s["kort"] for s in dom["sheets"]] ==
      [[dagar[di][1][i]["id"] for di, idx in sida for i in idx] for sida in plan], "")
    K("PDF: A4", not pk["a4"], ", ".join(pk["a4"]))
    K("PDF: antal sidor = planerat", pk["sidor"] == len(plan), f"{pk['sidor']} / {len(plan)}")
    href = {k["id"]: k["href"] for k in dom["kort"]}
    bad = []
    for i, (s, lankar) in enumerate(zip(dom["sheets"], pk["lankar"])):
        if sorted(href[x] for x in s["kort"]) != sorted(l["uri"] or "" for l in lankar):
            bad.append(f"sida {i + 1}: fel URI eller antal")
        bad += [f"sida {i + 1}: länkyta {l['bredd']:.0f} mm" for l in lankar if abs(l["bredd"] - 190) > 1.5]
    K("PDF: en länk per kort, ca 190 mm bred, rätt URI", not bad, ", ".join(bad))
    pmarg = [min((l["botten"] for l in lankar), default=99) for lankar in pk["lankar"]]
    bad = [f"sida {i + 1}: {v:.1f} mm" for i, v in enumerate(pmarg) if v < 9.8]
    K("PDF: sista kortet minst 10 mm ovanför underkanten", not bad, ", ".join(bad))
    ej = [n for n, inb in pk["typsnitt"].items() if not inb]
    K("PDF: typsnitten inbäddade", not ej, ", ".join(ej) or ", ".join(sorted(pk["typsnitt"])))
    if rk["metod"] is None:
        K("Raster: sidorna rastrerade", False,
          "varken pypdfium2 eller pdftoppm finns – pip install pypdfium2 pillow")
    elif any(h is None for h in rk["horn"]):
        K("Raster: sidorna rastrerade (hörnkontroll kräver pillow)", True, rk["metod"])
    else:
        bad = [f"sida {i + 1}" for i, ok in enumerate(rk["horn"]) if not ok]
        K("Raster: gradienten går ut i alla fyra hörn", not bad, ", ".join(bad))
    K("Mobil: ingen horisontell scroll vid 390 px", sw <= 390, f"scrollWidth {sw}")


def typsnitt_avvikelser(miljo: dict) -> list:
    t = miljo.get("typsnitt") or {}
    if not t or "fel" in t:
        return ["Typsnittet kunde inte avläsas i webbläsaren – kontrollera PDF:ens typsnitt."]
    familjer = {f for v in t.values() for f, _ in v}
    ut = []
    if familjer == {"Arial"}:
        pass
    elif familjer <= {"Arial", "Liberation Sans"}:
        ut.append("Arial saknas i exportmiljön; Liberation Sans (metriskt lika) användes.")
    else:
        ut.append(f"Varken Arial eller Liberation Sans fanns; {', '.join(sorted(familjer))} användes. "
                  "Installera fonts-liberation för rätt layout.")
    ps = " ".join(p for _, p in t.get(".sub b", []))
    if not re.search(r"black|heavy", ps, re.I):
        ut.append("Vikt 900 (regionnamnet i underraden) finns inte i typsnittet och visas som fet (700).")
    return ut


def skriv_rapport(m: dict, resultat: list, miljo: dict) -> str:
    rader = ["| Region | Verifierade evenemang | varav kompetensutveckling/företagande | PDF-sidor | Obekräftade/uteslutna |",
             "| --- | --- | --- | --- | --- |"]
    for r in resultat:
        reg = r["reg"]
        ob = sum(e["status"] == "obekraftad" for e in reg["B"])
        kb = sum(e.get('kategori') != 'Arbetsmarknad' for e in reg['A'])
        rader.append(f"| {reg['kort']} | {len(reg['A'])} | {kb} | {r['sidor'] or '–'} | {ob}/{len(reg['B']) - ob} |")
    rader += ["", "- **Osäkerheter:**"]
    poster = []
    for r in resultat:
        for e in r["reg"]["A"]:
            f = []
            if e["tidsfrist"] and e["tidsfrist"] < e["kd"]:
                f.append(f"anmälningsfristen {e['tidsfrist'].isoformat()} har passerat (kortet visar "
                         "'Anmälan krävs')")
            if e["status"] == "vantelista":
                f.append("fullbokat – väntelista")
            if e["prelim"]:
                f.append("preliminär lokal")
            if e["samlingssida"]:
                f.append("länken går till en samlingssida")
            if not e["pris_tagg"]:
                f.append(f"priset ({e['pris']}) visas inte på kortet" if e["pris"] else "pris ej angivet")
            if e["osakerheter"]:
                f.append(e["osakerheter"].rstrip("."))
            if f:
                titel = e["titel"].replace("­", "")
                poster.append(f"  - {e['id']} {titel} ({kortdatum(e['visning'])}): {'; '.join(f)}.")
    rader += poster or ["  - Inga."]
    rader.append("- **Avvikelser från specifikationen:**")
    av = typsnitt_avvikelser(miljo) + [f"Standardval: {v}" for v in m["val"]]
    rader += [f"  - {a}" for a in av] or ["  - Inga."]
    rader.append("- **Researchluckor:**")
    for r in resultat:
        reg = r["reg"]
        rader.append(f"  - {reg['kort']}: " + ("; ".join(reg["luckor"]) if reg["luckor"] else "inga") + ".")
    return "\n".join(rader) + "\n"


def bygg(data_path: Path, ut: Path, tillat_exempel: bool = False, tyst: bool = False) -> int:
    data = ladda_json(data_path)
    c, m = validera(data, tillat_exempel)
    visa_meddelanden(c)
    if m and not tyst:
        visa_tackning(m)
    if c.fel or not m:
        print(f"\nBYGGET AVBRÖTS: {len(c.fel)} fel i {data_path}. Rätta data.json och kör igen.")
        return 2
    saknas = [n for n in ("playwright", "openpyxl", "pypdf") if importlib.util.find_spec(n) is None]
    if saknas:
        fel_avbryt(f"paket saknas: {', '.join(saknas)}. "
                   "Installera: pip install playwright openpyxl pypdf pypdfium2 pillow")
    from playwright.sync_api import sync_playwright
    et = etiketter(m["period"])
    ut.mkdir(parents=True, exist_ok=True)
    qa = data_path.parent / "qa"
    qa.mkdir(parents=True, exist_ok=True)
    for monster in (f"{et['prefix']}_*_Evenemang.html", f"{et['prefix']}_*_Evenemang.pdf"):
        for gammal in ut.glob(monster):
            gammal.unlink()
    for gammal in qa.glob(f"{et['prefix']}_*.png"):
        gammal.unlink()
    xlsx = ut / f"{et['prefix']}_Evenemang_Källunderlag.xlsx"
    skriv_excel(xlsx, m, et)

    resultat, miljo = [], {}
    with sync_playwright() as pw:
        browser = starta_chromium(pw)
        miljo["chromium"] = browser.version
        for reg in m["regioner"]:
            res = {"reg": reg, "kontroller": [], "sidor": 0, "bilder": []}
            resultat.append(res)
            if reg["A"]:
                bygg_region(browser, reg, m, et, ut, qa, res, miljo)
        browser.close()

    rapport = skriv_rapport(m, resultat, miljo)
    (data_path.parent / "rapport.md").write_text(rapport, encoding="utf-8")
    kopia = ut / f"{et['prefix']}_data.json"
    if kopia.resolve() != data_path.resolve():
        shutil.copyfile(data_path, kopia)

    p = m["period"]
    print(f"\nPERIOD {et['namn']} {p['ar']}: {dagrubrik(p['start']).lower()}–{dagrubrik(p['slut']).lower()} "
          f"{p['slut'].year} · kontrolldatum {m['kd']} · Chromium {miljo['chromium']} · "
          f"raster {miljo.get('raster') or '–'}")
    antal_fel = 0
    for r in resultat:
        reg = r["reg"]
        print(f"\nKONTROLLER {reg['kort']} – {len(reg['A'])} kort, {r['sidor']} sidor")
        if not reg["A"]:
            print("  –    inga verifierade evenemang: ingen HTML/PDF skapad (står i Excel och rapport)")
        for namn, ok, detalj in r["kontroller"]:
            antal_fel += not ok
            print(f"  {'OK ' if ok else 'FEL'}  {namn}" + (f" – {detalj}" if detalj else ""))
        for e in reg["A"]:
            print(f"       {e['id']:<5} {kortdatum(e['visning']):<7} {e['tid_text']:<12} {e['titel'][:70]}")
    print("\nFILER")
    for f in sorted(ut.glob(f"{et['prefix']}_*")):
        print(f"  {f}")
    bilder = [b for r in resultat for b in r["bilder"]]
    if bilder:
        print("\nGRANSKA BILDERNA (levereras inte):")
        for b in bilder:
            print(f"  {b}")
    print(f"\nRAPPORT: {data_path.parent / 'rapport.md'}")
    if antal_fel:
        print(f"\n{antal_fel} KONTROLL(ER) MISSLYCKADES – rätta data.json och kör bygg igen.")
        return 1
    print("\nALLA KONTROLLER GODKÄNDA")
    return 0

# ---------------------------------------------------------------- kommandon


def ladda_json(path: Path):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except FileNotFoundError:
        fel_avbryt(f"hittar inte {path}")
    except json.JSONDecodeError as ex:
        fel_avbryt(f"{path} är inte giltig JSON (rad {ex.lineno}, kolumn {ex.colno}): {ex.msg}")


def skapa_data(regioner: list, typ: str, nr: int, ar: int, idag: dt.date, tidigare=None) -> dict:
    start, slut = berakna_period(typ, nr, ar)
    tid = manadsord(start, slut)
    data = {
        "version": VERSION,
        "period": {"typ": typ, "nr": nr, "ar": ar, "start": start.isoformat(), "slut": slut.isoformat()},
        "kontrolldatum": idag.isoformat(),
        "val": [],
        "_mall_evenemang": MALL_EVENEMANG,
        "_mall_kalla": MALL_KALLA,
        "regioner": [],
    }
    for r in regioner:
        gamla = []
        for g in (tidigare or {}).get("regioner", []):
            if g.get("namn") == r["namn"]:
                sedda = set()
                for k in g.get("kallor") or []:
                    url = txt(k.get("url"))
                    if url and k.get("typ") != "sokning" and url not in sedda:
                        sedda.add(url)
                        gamla.append({"typ": k.get("typ"), "kommun": k.get("kommun"), "url": url,
                                      "resultat_forra": k.get("resultat")})
        data["regioner"].append({
            "namn": r["namn"], "kort": r["kort"], "kod": r["kod"], "kommuner": r["kommuner"],
            "sokfragor": [f"{o} {r['kort']} {tid}" for o in SOKORD]
                         + [f"{o} {tid}" for o in SOKORD_NATIONELLT],
            "tidigare_kallor": gamla, "kallor": [], "researchluckor": [], "evenemang": []})
    return data


def cmd_init(a) -> int:
    idag = idag_sthlm()
    typ, nr = tolka_period(a.period)
    if a.ar in (None, "", "auto"):
        ar = auto_ar(typ, nr, idag)
    else:
        try:
            ar = int(a.ar)
        except ValueError:
            fel_avbryt("--ar ska vara ett årtal eller auto")
    try:
        start, slut = berakna_period(typ, nr, ar)
    except ValueError:
        fel_avbryt(f"vecka {nr} finns inte år {ar}")
    regioner = [tolka_region(r) for r in a.regioner]
    tilldela_koder(regioner)
    tidigare = ladda_json(Path(a.tidigare)) if a.tidigare else None
    ut = Path(a.ut)
    if ut.exists() and not a.skriv_over:
        fel_avbryt(f"{ut} finns redan – använd --skriv-over om den ska ersättas")
    ut.parent.mkdir(parents=True, exist_ok=True)
    data = skapa_data(regioner, typ, nr, ar, idag, tidigare)
    ut.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    et = etiketter({"typ": typ, "nr": nr, "ar": ar})
    print(f"Skapade {ut}")
    print(f"PERIOD   {et['namn']} {ar}: {dagrubrik(start).lower()} – {dagrubrik(slut).lower()}"
          f"  ({intervall(start, slut)})")
    print(f"KONTROLLDATUM {idag.isoformat()} (Europe/Stockholm)")
    for r, d in zip(regioner, data["regioner"]):
        print(f"\n{r['namn']} – ID-prefix {r['kod']} – {len(r['kommuner'])} kommuner:")
        print("  " + ", ".join(r["kommuner"]) if r["kommuner"] else "  (fyll i kommuner själv)")
        print(f"  {len(d['sokfragor'])} standardsökningar, t.ex. \"{d['sokfragor'][0]}\"")
        if d["tidigare_kallor"]:
            print(f"  {len(d['tidigare_kallor'])} källor från förra körningen att kontrollera först")
    return 0


def cmd_kontroll(a) -> int:
    c, m = validera(ladda_json(Path(a.data)))
    visa_meddelanden(c)
    if m:
        visa_tackning(m)
    if c.fel or not m:
        print(f"\n{len(c.fel)} FEL – rätta innan bygg.")
        return 2
    print(f"\nINGA FEL ({len(c.varn)} varningar).")
    return 0


def exempeldata(idag: dt.date) -> dict:
    """Fiktivt exempel för miljötestet (länkar till .invalid, får aldrig levereras)."""
    typ, nr = "vecka", (idag + dt.timedelta(days=7)).isocalendar()[1]
    ar = auto_ar(typ, nr, idag)
    data = skapa_data([tolka_region("Kalmar län (K)"), tolka_region("Gotlands län (G)")], typ, nr, ar, idag)
    s, _ = berakna_period(typ, nr, ar)
    D = lambda n: (s + dt.timedelta(days=n)).isoformat()
    kd = idag.isoformat()
    bas = {"arrangor": "Exempelkommunen, näringslivsenheten", "format": "plats", "kommun": "Kalmar",
           "ort": "Kalmar", "anmalan": "kravs", "pris_tagg": "Gratis", "kalla": "Fiktivt testfall",
           "belagg": "", "tid_slut": None, "status": "verifierad", "lank": "https://exempel.invalid/",
           "kategori": "Arbetsmarknad"}
    def ev(**k):
        e = dict(bas, **k)
        e["lank"] = e["lank"] + str(abs(hash(e["titel"])) % 10000)
        d = dt.date.fromisoformat(e["datum"])
        e["belagg"] = e["belagg"] or f"{dagrubrik(d)} {d.year}, kl. {e['tid_start']}"
        return e
    data["regioner"][0]["evenemang"] = [
        ev(titel="Frukostmöte: Så hittar du rätt kompetens", kategori="Företagande", typ="Frukostmöte", datum=D(0), tid_start="07:30",
           tid_slut="09:00", lokal="Stadshuset", adress="Storgatan 1",
           arrangor="Exempelkommunen, näringslivsenheten i samarbete med Almi Exempel",
           beskrivning="Frukostmöte om rekrytering och kompetensförsörjning för små och medelstora företag.",
           tidsfrist=D(-1) if s > idag else kd),
        ev(titel="Rekryteringsträff lager och logistik", typ="Rekryteringsträff", datum=D(0), tid_start="10:00",
           tid_slut="12:00", kommun="Nybro", ort="Nybro", lokal="Arbetsförmedlingen", anmalan="ingen",
           villkor=["För arbetssökande"], arrangor="Arbetsförmedlingen Exempel",
           beskrivning="Rekryteringsträff där lokala arbetsgivare inom lager och logistik möter arbetssökande."),
        ev(titel="Yrkesmässan Sydost", typ="Mässa", datum=D(1), datum_slut=D(3), tid_start="09:00",
           tid_slut="15:00", lokal="Exempelhallen", adress="Hamngatan 12", pris_tagg="Fri entré", anmalan="ingen",
           beskrivning="Mässa för dig som vill byta yrke eller studera vidare, med utställare från hela länet."),
        ev(titel="Webbinarium: Starta eget i Kalmar län", kategori="Företagande", typ="Seminarium", datum=D(1), tid_start="12:00",
           tid_slut="13:00", format="digitalt", kommun=None, ort=None, digital_plattform="Teams",
           regional_koppling="Arrangör och målgrupp i Kalmar län", tidsfrist=(idag - dt.timedelta(days=2)).isoformat(),
           beskrivning="Seminarium för blivande företagare om affärsplan, bolagsform och första kunderna."),
        ev(titel="Öppet hus på Vuxenutbildningen", typ="Öppet hus", datum=D(2), tid_start="15:00",
           tid_slut="18:00", dropin="15:00-18:00", kommun="Oskarshamn", ort="Oskarshamn", lokal="Komvux",
           anmalan="ingen", beskrivning="Öppet hus för vuxna som vill läsa in gymnasiet, byta yrke eller lära sig svenska."),
        ev(titel="Nätverksträff för kvinnliga företagare", typ="Nätverksträff", datum=D(2), tid_start="17:30",
           tid_slut="20:00", kommun="Västervik", ort="Västervik", lokal="Exempelhotellet", preliminar_lokal=True,
           pris_tagg="150 kr medlem / 250 kr övriga ex moms", villkor=["Begränsat antal platser", "Endast kvinnor"],
           tidsfrist=D(1), beskrivning="Nätverksträff för kvinnor som driver eller vill starta företag i norra länet."),
        ev(titel="Lunchträff HR-nätverket", typ="Lunchträff", datum=D(2), tid_start="11:30", tid_slut="13:00",
           status="vantelista", lokal="Exempelhuset", tidsfrist=D(0),
           beskrivning="Lunchträff för HR-ansvariga om kompetensbaserad rekrytering och introduktion."),
        ev(titel="Branschdag besöksnäring: framtidens kompetensbehov", typ="Konferens", datum=D(3),
           tid_start="09:00", tid_slut="16:00", kommun="Borgholm", ort="Borgholm", lokal="Exempelgården",
           pris_tagg="1200 kr ex moms", villkor=["Uteblivande 500 kr ex moms", "Bindande anmälan"],
           ovrigt="Lunch ingår", tidsfrist=D(1),
           beskrivning="Konferens om rekrytering, utbildning och säsongsarbete inom turism och besöksnäring."),
        ev(titel="Workshop: Digital marknadsföring för småföretag", kategori="Kompetensutveckling", typ="Workshop", datum=D(3), tid_start="13:00",
           tid_slut="16:00", format="hybrid", digital_plattform="Zoom", kommun="Vimmerby", ort="Vimmerby",
           lokal="Exempelcenter", pris_tagg=None, samlingssida=True,
           beskrivning="Workshop för småföretagare om annonsering i sociala medier och mätning av resultat."),
        ev(titel="Jobbspanarträff", typ="Rekryteringsträff", datum=D(4), tid_start="09:00", kommun="Emmaboda",
           ort="Emmaboda", lokal="Biblioteket", anmalan="okand",
           beskrivning="Rekryteringsträff för arbetssökande som vill hitta lediga jobb och träffa arbetsgivare."),
        ev(titel="Entreprenörskväll med pitchtävling", kategori="Företagande", typ="Nätverksträff", datum=D(4), tid_start="18:00",
           tid_slut="21:00", lokal="Exempelscenen", villkor=["Från 18 år"], ovrigt="Del 2: workshop 23 okt",
           beskrivning="Nätverksträff för entreprenörer och investerare, med pitchar från nystartade bolag."),
    ]
    data["regioner"][0]["evenemang"] += [
        ev(titel="Jobbmässa Kalmar", status="obekraftad", datum=D(4), tid_start=None,
           skal="Endast aggregator; arrangörens sida visar förra årets mässa", beskrivning=None),
        ev(titel="Karriärdag för studenter", status="utesluten", regel="R3", datum=D(2), tid_start="10:00",
           skal="Endast studenter vid lärosätet", beskrivning=None),
        ev(titel="Företagsfrukost Torsås", status="installt", datum=D(1), tid_start="07:30", kommun="Torsås",
           ort="Torsås", skal="Inställt enligt arrangören", beskrivning=None),
    ]
    data["regioner"][0]["kallor"] = [
        {"typ": "kommun", "kommun": "Kalmar", "url": "https://exempel.invalid/kalmar", "resultat": "kandidater"},
        {"typ": "komvux", "kommun": "Kalmar", "url": "https://exempel.invalid/komvux", "resultat": "inga"},
        {"typ": "af", "kommun": None, "url": "https://exempel.invalid/af", "resultat": "kandidater"},
        {"typ": "region", "kommun": None, "url": "https://exempel.invalid/region", "resultat": "ej_lasbar"},
        {"typ": "sokning", "fraga": data["regioner"][0]["sokfragor"][0], "resultat": "kandidater"},
    ]
    data["regioner"][1]["evenemang"] = [
        ev(titel="Medlemsfrukost", status="utesluten", regel="R4", datum=D(1), tid_start="07:30",
           kommun="Gotland", ort="Visby", skal="Endast för medlemmar", beskrivning=None),
    ]
    data["val"] = ["Testkörning med fiktiva data."]
    return data


def cmd_test(a) -> int:
    print(f"bygg_kalender.py {VERSION} · Python {sys.version.split()[0]}")
    assert sum(len(k.split(",")) for *_, k in LAN) == 290, "kommunlistan ska ha 290 kommuner"
    tmp = Path(tempfile.mkdtemp(prefix="evkal_test_"))
    try:
        data_p = tmp / "arbete" / "data.json"
        data_p.parent.mkdir(parents=True)
        data_p.write_text(json.dumps(exempeldata(idag_sthlm()), ensure_ascii=False, indent=2), encoding="utf-8")
        kod = bygg(data_p, tmp / "leverans", tillat_exempel=True, tyst=True)
        print("\n" + (data_p.parent / "rapport.md").read_text(encoding="utf-8"))
        print("MILJÖTEST " + ("GODKÄNT" if kod == 0 else "MISSLYCKADES"))
        if a.behall:
            print(f"Testfilerna finns kvar i {tmp}")
        return kod
    finally:
        if not a.behall:
            shutil.rmtree(tmp, ignore_errors=True)


def main(argv=None) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ap = argparse.ArgumentParser(prog="bygg_kalender.py", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--version", action="version", version=VERSION)
    sub = ap.add_subparsers(dest="kommando", required=True)
    t = sub.add_parser("test", help="kontrollera miljön med ett fiktivt exempel")
    t.add_argument("--behall", action="store_true", help="behåll testfilerna")
    i = sub.add_parser("init", help="skapa data.json")
    i.add_argument("--regioner", nargs="+", required=True, help='t.ex. "Kalmar län (K)" "Skåne län (S)"')
    i.add_argument("--period", required=True, help='"vecka 42" eller "månad oktober"')
    i.add_argument("--ar", default="auto", help="årtal eller auto (standard)")
    i.add_argument("--tidigare", help="data.json från förra körningen (källorna kontrolleras först)")
    i.add_argument("--ut", default="arbete/data.json")
    i.add_argument("--skriv-over", action="store_true")
    k = sub.add_parser("kontroll", help="validera data.json och visa täckningen")
    k.add_argument("data")
    b = sub.add_parser("bygg", help="bygg Excel, HTML, PDF och kör kontrollerna")
    b.add_argument("data")
    b.add_argument("--ut", default="leverans")
    a = ap.parse_args(argv)
    if a.kommando == "test":
        return cmd_test(a)
    if a.kommando == "init":
        return cmd_init(a)
    if a.kommando == "kontroll":
        return cmd_kontroll(a)
    return bygg(Path(a.data), Path(a.ut))


if __name__ == "__main__":
    sys.exit(main())
